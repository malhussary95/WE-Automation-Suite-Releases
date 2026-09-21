import os
import re
from datetime import datetime
from pathlib import Path
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
import platform
import subprocess

import openpyxl
import pandas as pd
from .field_handlers import apply_selected_fields
from .domain import extractors
from .loaders import (
    build_resource_result,
    enrich_installed_base_if_needed,
    load_branch_json,
    load_customer_json,
    load_installed_base_data,
    load_resource_data,
    load_commercial_data,
    load_service_order_data,
    load_task_data,
    load_user_json,
    resolve_order_from_cid,
)
from .cancellation import ExtractionCancelled, check_cancelled
from .ui_bridge import reset_ui, set_progress, show_error, show_info

from ecrm_extractor.core.storage import save_resource_json, save_json
from ecrm_extractor.fields import (
    FIELD_OPTIONS,
    SO_REQUIRED_FIELDS,
    RESOURCE_REQUIRED_FIELDS,
    RESULT_REQUIRED_FIELDS,
    IB_REQUIRED_FIELDS,
    CUSTOMER_JSON_REQUIRED_FIELDS,
    USER_JSON_REQUIRED_FIELDS,
    BRANCH_JSON_REQUIRED_FIELDS,
)

from ecrm_extractor.utils.helpers import has_any_selected


def run_extraction(deps, gui):
    """المدخل الرئيسي لعملية الاستخراج مع معالجة الأخطاء العامة"""
    try:
        _run_extraction(deps, gui)

    except ExtractionCancelled:
        # Cancellation is a normal user action. Keep the partial Excel file
        # and keep the real progress instead of resetting it to zero.
        saved_count = int(getattr(gui, "_partial_results_count", 0) or 0)
        progress = float(getattr(gui, "_last_progress", 0) or 0)
        reset_ui(gui, reset_progress=False)
        set_progress(gui, progress, f"CANCELLED - Saved {saved_count} item(s)")

    except Exception as exc:
        set_progress(gui, 0, "FAILED")
        reset_ui(gui)
        
        # محاولة إظهار رسالة خطأ مفصلة
        show_error(
            gui,
            "Error",
            f"Extraction failed\n{exc}"
        )

    finally:
        # التأكد من إغلاق متصفح البورتال عند الانتهاء لتوفير موارد الجهاز
        if hasattr(deps, "cleanup_ftth"):
            deps.cleanup_ftth()

def _web_stage(gui, message, phase=None):
    """Emit a detailed stage message when running through the web adapter."""
    try:
        if phase and hasattr(gui, "set_web_phase"):
            gui.set_web_phase(phase)
        if hasattr(gui, "status_var"):
            gui.status_var.set(message)
    except Exception:
        pass

def _run_extraction(deps, gui):
    # 1. إعدادات البداية
    _web_stage(gui, "Reading input and validating extraction request…", "Reading Input")
    debug_mode = getattr(gui, "debug_mode", False)
    direct_values = getattr(gui, "direct_values", None) or []
    if not gui.file_path and not direct_values:
        return show_error(gui, "Error", "Choose Excel File or paste values")

    # 2. قراءة البيانات وتحديد الأوامر المطلوبة
    df = _read_input_data(gui, direct_values)
    if df is None: return
    _web_stage(gui, f"Input loaded: {len(df)} row(s)", "Resolving Items")
    check_cancelled(gui)

    items = _resolve_items_by_mode(deps, gui, df, debug_mode)
    if not items:
        return show_error(gui, "Error", "No valid items found to process")
    _web_stage(gui, f"Ready: {len(items)} item(s) to process", "Creating Session")
    check_cancelled(gui)

    # 3. حفظ بيانات الاعتماد وتحضير ملف المخرجات
    deps.save_credentials(gui.username.strip(), gui.password.strip())
    _web_stage(gui, "Creating authenticated ECRM session…", "Creating Session")
    session = deps.create_session()
    _web_stage(gui, "ECRM session created successfully.", "Processing")
    selected = set(gui.selected_fields)
    
    out_excel, headers, wb, ws = _prepare_output_excel(selected, gui.mode)
    print(f"DEBUG: Output file initialized at: {out_excel}")

    # استدعاء تصحيح أخطاء CID إذا كان الوضع مختاراً
    if gui.mode == "CID" and debug_mode:
        cid_values = [it.get("cid") for it in items if isinstance(it, dict)]
        _debug_cid_mode(deps, gui, session, cid_values, selected, debug_mode)
    
    # 4. حلقة المعالجة الرئيسية
    try:
        _web_stage(gui, f"Processing {len(items)} item(s)…", "Processing")
        processed_rows = _process_items_loop(deps, gui, session, items, selected, out_excel, headers, wb, ws, debug_mode)
        set_progress(gui, 100, "DONE")
        _publish_smart_result(gui, processed_rows, out_excel, headers)
        show_info(gui, "Completed", f"Saved successfully to:\n{out_excel}")
        _open_output_folder(out_excel)
    except Exception as e:
        # حتى لو فشل الاستخراج، نخبر المستخدم بمكان الملف الذي قد يحتوي على بيانات جزئية
        show_error(gui, "Partial Error", f"An error occurred, but partial data might be saved at:\n{out_excel}\n\nError: {e}")
        raise e

def _read_input_data(gui, direct_values):
    if direct_values:
        columns_by_mode = {
            "ORDER": "Order",
            "CID": "CID",
            "SO": "SO",
            "ORD": "ORD",
        }
        column = columns_by_mode.get(gui.mode, "Order")
        return pd.DataFrame({column: direct_values})

    try:
        return pd.read_excel(gui.file_path, dtype={"Order": str, "CID": str, "SO": str, "ORD": str})
    except Exception as exc:
        show_error(gui, "Error", f"Cannot open Excel\n{exc}")
        return None

def _resolve_items_by_mode(deps, gui, df, debug_mode):
    """تحديد قائمة الـ Orders بناءً على وضع الاستخراج المختارة"""
    mode = gui.mode
    session = deps.create_session()
    selected = set(gui.selected_fields)

    if mode == "CID":
        if "CID" not in df.columns: return None
        cid_values = list(
            dict.fromkeys(
                cid
                for cid in df["CID"].dropna().astype(str).str.strip().tolist()
                if cid
            )
        )
        set_progress(gui, 0, f"Ready to process {len(cid_values)} CIDs")
        return [{"cid": cid} for cid in cid_values]

    elif mode == "SO":
        if "SO" not in df.columns: return None
        so_values = df["SO"].dropna().astype(str).str.strip().tolist()
        items = []
        for so_id in so_values:
            check_cancelled(gui)
            service_orders = deps.get_service_order_by_soid(session, so_id)
            if service_orders:
                so = service_orders[0]
                if so.get("te_orderlinenumber"): items.append(str(so["te_orderlinenumber"]))
                if debug_mode: save_json(so_id, "SERVICE_ORDER_SEARCH", service_orders)
        return items

    elif mode == "ORD":
        if "ORD" not in df.columns: return None
        ord_values = df["ORD"].dropna().astype(str).str.strip().tolist()
        items = []
        for ord_num in ord_values:
            check_cancelled(gui)
            res = deps.get_resources_by_ord(session, ord_num)
            if res:
                if res[0].get("new_orderlinenumberid"): items.append(str(res[0]["new_orderlinenumberid"]))
                if debug_mode: save_json(ord_num, "ORD_SEARCH", res)
        return items

    else: # ORDER MODE
        if "Order" not in df.columns: return None
        return df["Order"].dropna().astype(str).str.strip().tolist()

def _debug_cid_mode(deps, gui, session, cid_values, selected, debug_mode):
    if not debug_mode:
        return
    for cid in cid_values:
        try:
            check_cancelled(gui)
            resources = deps.get_resource_by_cid(session, cid)
            save_resource_json("CID_DEBUG", cid, resources)
            # تحسين: لجلب بيانات التصحيح الكاملة، نكتفي بجلب بيانات المورد الأول فقط
            # لتقليل عدد استدعاءات الـ API في وضع التصحيح.
            if resources and resources[0].get("te_resourceregistryid"):
                first_resource_id = resources[0]["te_resourceregistryid"]
                full_resource_data = deps.get_resource_full_data(session, first_resource_id)
                if full_resource_data: save_resource_json("CID_FULL_DEBUG", cid, [full_resource_data])
        except ExtractionCancelled:
            raise
        except: pass

def _open_output_folder(file_path):
    """فتح المجلد الذي يحتوي على الملف المستخرج تلقائياً"""
    try:
        path = os.path.realpath(os.path.dirname(file_path))
        if platform.system() == "Windows":
            os.startfile(path)
        elif platform.system() == "Darwin":  # macOS
            subprocess.Popen(["open", path])
        else:  # linux variants
            subprocess.Popen(["xdg-open", path])
    except Exception as e:
        print(f"Could not open folder: {e}")

def _publish_smart_result(gui, rows, out_excel, headers):
    callback = getattr(gui, "smart_result_callback", None)
    if not callback:
        return

    try:
        callback(rows, out_excel, headers)
    except Exception as exc:
        print(f"Could not publish smart chat result: {exc}")

def _prepare_output_excel(selected, mode):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    # Determine output folder relative to script location for better portability
    # استخدام .resolve() يضمن الحصول على المسار الكامل والحقيقي
    base_dir = Path(__file__).resolve().parent.parent
    output_folder = base_dir / "output"
    output_folder.mkdir(parents=True, exist_ok=True)
    out_excel = (output_folder / f"ECRM_OUTPUT_{timestamp}.xlsx").resolve()
    
    headers = []
    if mode == "CID":
        headers.append("CID")
    
    # إضافة رقم الطلب في البداية إذا تم اختياره
    if "Order" in selected:
        headers.append("Order")

    # استثناء الفئات وإضافة الحقول المختارة
    categories = {"CID", "Order", "Request Number", "ESPT & infra status", "Network Data", "MSAN Data"}
    for field in FIELD_OPTIONS:
        if field not in categories and field in selected:
            headers.append(field)

    # توسيع الحقول التلقائية
    if "Network Data" in selected:
        headers.extend(["VLAN", "Interface", "PE IP", "WAN IP", "PE Name", "VRF", "LAN", "Description", "Config File"])
    if "MSAN Data" in selected:
        headers.extend(["MSAN Code", "MSAN IP", "MSAN Name", "Shelf", "Card", "Port", "Comment", "Description", "RR Port"])
    if "Work Order PDF" in selected or "ESPT & infra status" in selected:
        headers.append("WO issue Date")
        if "Work Order PDF" not in headers: headers.append("Work Order PDF")
        if "Work Order PDF" in selected:
            headers.extend(["Cabinet", "Splitter", "OLT Name", "Slot", "Port", "ONU Serial Number", "CST Name"])
    if "CID" in selected:
        headers.extend(["Old Circuit ID", "New Circuit ID"])
    if "Request Number" in selected or "ONU Tech Data" in selected:
        headers.extend(["Old ORD", "New ORD"])
    if "ESPT & infra status" in selected:
        headers.extend(["Infra status", "ESPT status"])

    if "ONU Tech Data" in selected:
        headers.extend([
            "MSAN IP",
            "MSAN Code",
            "Port",
            "ONT ID",
            "ONU Serial Number",
            "ONU Vendor",
            "Cabinet vendor",
            "ISPName",
            "Down Speed",
            "Acceptance Date"
        ])

    if "Latest Migration by E-Support SO" in selected:
        headers.extend(["Migration SO Type", "Migration SO Status", "Migration Current Task", "Migration Current Task Owner"])

    headers = list(dict.fromkeys(headers)) # حذف التكرار (يجب أن يكون في النهاية)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Results"
    ws.append(headers)
    return str(out_excel), headers, wb, ws

def _process_items_loop(deps, gui, session, items, selected, out_excel, headers, wb, ws, debug_mode):
    # NOTE: This function was partially modified to preserve input ordering.
    # Final Excel writing should be done after all futures complete in index order.

    total = len(items)
    shared_state = {"wo_pdf_decision": None}
    lock = threading.Lock()
    completed_count = 0
    processed_rows = []

    # ترتيب output بنفس ترتيب input (مع parallel workers)
    results_by_index = [None] * total

    # Dedicated rows for CID_Status: one row per CID.
    cid_status_by_index = [None] * total

    # Configurable I/O concurrency. ECRM extraction is network/API-bound, so
    # multiple workers can reduce total runtime substantially. Each worker
    # owns its own requests.Session (see thread-local session below).
    try:
        requested_workers = int(getattr(gui, "max_workers", 5) or 5)
    except (TypeError, ValueError):
        requested_workers = 5
    max_workers = max(1, min(requested_workers, 12))
    max_workers = min(max_workers, total) if total else 1
    stage_msg = getattr(gui, "_emit", None)
    if stage_msg:
        try:
            stage_msg(f"Parallel workers: {max_workers}", "info")
        except Exception:
            pass

    # Reuse one authenticated session per worker thread instead of creating a
    # new NTLM session for every item. Session creation/handshake is expensive
    # and was a significant source of avoidable latency.
    thread_local = threading.local()
    
    def worker(index, item):
        nonlocal completed_count
        # Each thread needs its own session to avoid NTLM state conflicts
        emit = getattr(gui, "_emit", None)
        def stage(message, phase="Processing"):
            try:
                if hasattr(gui, "set_web_phase"):
                    gui.set_web_phase(phase)
                if emit:
                    emit(message, "info")
            except Exception:
                pass

        if not hasattr(thread_local, "session"):
            stage(f"Worker session {threading.current_thread().name}: creating authenticated session…", "Creating Session")
            thread_local.session = deps.create_session()
        thread_session = thread_local.session
        
        try:
            check_cancelled(gui)
            input_cid = ""
            if isinstance(item, dict):
                input_cid = str(item.get("cid") or "").strip()
                stage(f"Item {index + 1}/{total}: resolving CID {input_cid} to Order…", "Resolving CID")
                order = resolve_order_from_cid(thread_session, input_cid, deps)
                if not order:
                    # حفظ بيانات التصحيح حتى لو لم يتم العثور على الطلب
                    if debug_mode:
                        res = deps.get_resource_by_cid(thread_session, input_cid)
                        save_resource_json("CID_NOT_FOUND_DEBUG", input_cid, res)
                    
                    order = "" # لا نخرج، بل نكمل المحاولة باستخدام الـ CID
            else:
                order = str(item).strip()

            check_cancelled(gui)
            
            # Load data using the thread-specific session
            # التصحيح: إذا نجحنا في الوصول لرقم الطلب (Order) حتى لو كان الدخل CID، 
            # يجب استخدامه كمعرف أساسي لجلب الموارد لضمان استخراج الـ Request Number (Old/New ORD)
            if order:
                technical_id = order
                technical_mode = "ORDER"
            else:
                technical_id = input_cid
                technical_mode = "CID"

            stage(f"Item {index + 1}/{total}: loading Resource Registry ({technical_mode}={technical_id})…", "Resource Registry")
            resources = load_resource_data(thread_session, technical_id, selected, deps, debug=debug_mode)
            stage(f"Item {index + 1}/{total}: Resource Registry loaded ({len(resources or [])} record(s))", "Resource Registry")

            # Circuit Status is resolved directly by CID from Resource Registry.
            # This is separate from Order lookup because one CID may have
            # multiple Resource records.
            cid_resources = []
            if "Circuit Status" in selected and input_cid:
                try:
                    cid_resources = deps.get_resources_by_cids(
                        thread_session, [input_cid]
                    ) or []
                except Exception:
                    cid_resources = []

            
            # ترتيب الموارد تنازلياً لضمان أن كافة العمليات التالية (PDF, MSAN, ONU) تعمل على "الأحدث"
            if resources:
                try:
                    resources = sorted(
                        resources,
                        key=lambda x: x.get("createdon") or "",
                        reverse=True
                    )
                except:
                    pass
            
            # محاولة استخراج رقم الطلب من الموارد إذا لم يتم العثور عليه مسبقاً (خاص بوضع الـ CID)
            if not order and resources:
                for r in resources:
                    # البحث عن الحقول التي قد تحتوي على رقم الطلب في الموارد
                    val = (
                        r.get("te_requestnumber") or 
                        r.get("new_orderlinenumberid") or 
                        r.get("rr.te_requestnumber") or
                        r.get("_te_serviceorderid_value@OData.Community.Display.V1.FormattedValue")
                    )
                    if val:
                        val_str = str(val).strip()
                        if "ORD-" in val_str.upper() or (val_str.isdigit() and len(val_str) > 5):
                            order = val_str
                            break

            # الآن نقوم بجلب البيانات المعتمدة على رقم الطلب بعد تأكيد وجوده من الموارد
            if order:
                stage(f"Item {index + 1}/{total}: loading Service Order data…", "Service Order")
                service_orders, so, data_so, migration_so = load_service_order_data(
                    thread_session, order, "ORDER", selected, deps, debug=debug_mode
                )
                stage(f"Item {index + 1}/{total}: Service Order loaded", "Service Order")
            else:
                service_orders, so, data_so, migration_so = [], {}, {}, {}

            # Only fetch datasets required by the selected fields. The old web
            # adapter fetched Commercial/Tasks/Installed/Customer/User/Branch
            # data for every extraction, even when none of those fields were
            # selected. That could make a simple 3-field job look frozen.
            need_commercial = "Order Status" in selected
            need_tasks = bool(selected & {"Current Task", "Latest Migration by E-Support SO", "Migration Current Task", "Migration Current Task Owner"})
            need_installed = bool(selected & {"CST Name", "Branch", "Speed", "CST Name Arabic", "Branch Address", "Installed Resources", "CST Type", "CST Category", "CST Number", "Account manager", "Account manager mail"})
            need_customer = bool(selected & {"CST Name Arabic", "CST Type", "CST Category", "CST Number", "Account manager", "Account manager mail"})
            need_user = bool(selected & {"Account manager", "Account manager mail"})
            need_branch = "Branch Address" in selected

            if order and need_commercial:
                stage(f"Item {index + 1}/{total}: loading Commercial Order Lines…", "Commercial Orders")
                commercial_lines = load_commercial_data(thread_session, order, selected, deps, debug=debug_mode)
                stage(f"Item {index + 1}/{total}: Commercial Order Lines loaded ({len(commercial_lines or [])})", "Commercial Orders")
            else:
                commercial_lines = []

            if order and need_tasks:
                stage(f"Item {index + 1}/{total}: loading Service Order tasks…", "Tasks")
                tasks = load_task_data(thread_session, so, "ORDER", selected, deps)
                stage(f"Item {index + 1}/{total}: Tasks loaded ({len(tasks or [])})", "Tasks")
            else:
                tasks = []

            if order and need_installed:
                stage(f"Item {index + 1}/{total}: loading Installed Base…", "Installed Base")
                installed_raw = load_installed_base_data(thread_session, order, selected, deps, debug=debug_mode)
                installed = enrich_installed_base_if_needed(thread_session, installed_raw, order, selected, deps, debug=debug_mode)
                stage(f"Item {index + 1}/{total}: Installed Base loaded ({len(installed or [])})", "Installed Base")
            else:
                installed = []

            stage(f"Item {index + 1}/{total}: building selected resource fields…", "Building Result")
            result = build_resource_result(thread_session, resources, so, migration_so, technical_id, technical_mode, selected, deps, debug=debug_mode)

            l3 = extractors.extract_l3_data(resources, order=(order or input_cid), debug=debug_mode) if "Network Data" in selected else {}
            msan = extractors.extract_msan_data(resources) if "MSAN Data" in selected else {}

            if need_customer:
                stage(f"Item {index + 1}/{total}: loading customer data…", "Customer Data")
                customer_json = load_customer_json(thread_session, order, so, installed, selected, deps, debug=debug_mode)
            else:
                customer_json = {}

            if need_user:
                stage(f"Item {index + 1}/{total}: loading user/account data…", "User Data")
                user_json = load_user_json(thread_session, order, customer_json, selected, deps, debug=debug_mode)
            else:
                user_json = {}

            if need_branch:
                stage(f"Item {index + 1}/{total}: loading branch data…", "Branch Data")
                branch_data = load_branch_json(thread_session, order, so, installed, selected, deps, debug=debug_mode)
            else:
                branch_data = {}
            
            context = {
                "gui": gui, "debug_mode": debug_mode, "shared_state": shared_state,
                "deps": deps, "session": thread_session, "selected": selected, "order": order, "input_cid": input_cid,
                "service_orders": service_orders, "so": so, "migration_so": migration_so,
                "tasks": tasks, "resources": resources, "installed": installed,
                # Commercial Orders are loaded above and passed to
                # add_order_status() as the source of truth for Order Status.
                "commercial_lines": commercial_lines, "result": result, "l3": l3, "msan": msan,
                "customer_json": customer_json, "user_json": user_json, "branch_data": branch_data,
                "cid_resources": cid_resources
            }

            stage(f"Item {index + 1}/{total}: applying selected fields…", "Formatting Result")
            row = {}
            apply_selected_fields(row, context, selected)
            if input_cid: row["CID"] = input_cid


            # ============================================================
            # CID_Status: one row for every Old/New CID of the Order
            #
            # IMPORTANT:
            # "ESPT & infra status" is a normal result field and must NOT
            # trigger the CID_Status sheet or the per-CID Resource Registry
            # lookups below. Triggering this block for ESPT/Infra caused
            # unnecessary CID processing and slowed the extraction.
            # ============================================================
            cid_status_rows = []
            if (
                "Circuit Status" in selected
                or "Request Number" in selected
            ):
                circuit_ids = []
                if technical_mode == "ORDER":
                    old_ids = str(result.get("Old Circuit ID") or "").strip()
                    new_ids = str(result.get("New Circuit ID") or "").strip()

                    for raw_ids in (old_ids, new_ids):
                        if raw_ids:
                            for cid in re.split(r"\s*-\s*", raw_ids):
                                cid = str(cid).strip()

                                # Ignore invalid / empty CID values
                                if not cid:
                                    continue

                                if cid.lower() in {
                                    "none",
                                    "null",
                                    "nan",
                                    "n/a",
                                    "na",
                                    "-",
                                }:
                                    continue

                                if cid not in circuit_ids:
                                    circuit_ids.append(cid)
                elif input_cid:
                    circuit_ids.append(input_cid)

                # Always use the Order Status already calculated by
                # add_order_status(). Do NOT recalculate it here from
                # Installed Base or from a different Commercial Order rule.
                #
                # Order Status source/rules:
                #   - Commercial Order / Service only
                #   - newest first
                #   - Fulfilled -> Active
                #   - InFlight/Canceled -> check previous Service
                #   - Sales Termination -> Terminated
                #   - unknown/non-matching Order -> not found
                order_status = str(
                    row.get("Order Status") or "not found"
                ).strip()

                if not order_status:
                    order_status = "not found"

                for cid in circuit_ids:
                    circuit_status = "not found"
                    cid_ord = ""
                    infra_status = ""
                    espt_status = ""
                    te_note = ""

                    try:
                        # Always query Resource Registry by the CID itself.
                        # This is the source of the ORD for this exact CID,
                        # regardless of whether the input was Order, CID, SO,
                        # or ORD.
                        rr_list = deps.get_resources_by_cids(
                            thread_session, [cid]
                        ) or []
                    except Exception:
                        rr_list = []

                    if rr_list:
                        # ---------------------------------------------
                        # TE NOTE: collect all non-empty te_notes for this CID
                        # and join unique values with " - "
                        # ---------------------------------------------
                        notes = []
                        for rr in rr_list:
                            note = rr.get("te_notes")
                            if note is None:
                                continue
                            note = str(note).strip()
                            if not note or note.lower() in {"null", "none", "nan"}:
                                continue
                            if note not in notes:
                                notes.append(note)
                        te_note = " - ".join(notes)

                        # ---------------------------------------------
                        # CIRCUIT STATUS
                        # ---------------------------------------------
                        latest_rr = max(
                            rr_list,
                            key=lambda r: (
                                r.get("modifiedon")
                                or r.get("createdon")
                                or ""
                            )
                        )

                        def _rr_status_value(resource, field):
                            suffix = "@OData.Community.Display.V1.FormattedValue"
                            candidates = (
                                f"{field}{suffix}",
                                f"rr.{field}{suffix}",
                                f"rr_x002e_{field}{suffix}",
                                field,
                                f"rr.{field}",
                                f"rr_x002e_{field}",
                            )
                            for key in candidates:
                                value = resource.get(key)
                                if value not in (None, ""):
                                    return str(value).strip()
                            return ""

                        # CID lookup returns Resource Registry rows. These fields
                        # are explicitly requested by api_client.get_resources_by_cids().
                        # Prefer the latest RR, but fall back to older RR records
                        # for the same CID when the latest row has no value.
                        infra_status = _rr_status_value(
                            latest_rr, "te_infrastructureavailabilitystatusnewcode"
                        )
                        espt_status = _rr_status_value(
                            latest_rr, "te_esptstatuscode"
                        )

                        if not infra_status:
                            for rr in sorted(
                                rr_list,
                                key=lambda r: r.get("modifiedon") or r.get("createdon") or "",
                                reverse=True,
                            ):
                                infra_status = _rr_status_value(
                                    rr, "te_infrastructureavailabilitystatusnewcode"
                                )
                                if infra_status:
                                    break

                        if not espt_status:
                            for rr in sorted(
                                rr_list,
                                key=lambda r: r.get("modifiedon") or r.get("createdon") or "",
                                reverse=True,
                            ):
                                espt_status = _rr_status_value(rr, "te_esptstatuscode")
                                if espt_status:
                                    break

                        statecode = latest_rr.get("statecode")
                        if statecode == 0:
                            circuit_status = "Active"
                        elif statecode == 1:
                            circuit_status = "Inactive"
                        else:
                            statuscode = latest_rr.get("statuscode")
                            if statuscode == 1:
                                circuit_status = "Active"
                            elif statuscode == 2:
                                circuit_status = "Inactive"

                        # ---------------------------------------------
                        # ORD / REQUEST NUMBER FOR THIS CID
                        # ---------------------------------------------
                        # Do NOT use the input ORD here.
                        # We need the ORD belonging to the current CID.
                        #
                        # Prefer the newest Resource Registry record that
                        # actually contains te_requestnumber.
                        # The CID lookup endpoint may return a reduced
                        # Resource Registry payload. Therefore use BOTH:
                        #   1) the direct CID lookup result
                        #   2) the full "resources" payload already loaded
                        #      for this Order (this is the same source used
                        #      by All Results to build Old/New ORD).
                        rr_candidates = list(rr_list)

                        for resource in (resources or []):
                            resource_cid = (
                                resource.get("te_circuitid")
                                or resource.get("rr.te_circuitid")
                                or resource.get("rr_x002e_te_circuitid")
                                or ""
                            )
                            if str(resource_cid).strip() == cid:
                                rr_candidates.append(resource)

                        # Remove duplicate object references while preserving
                        # all available records.
                        unique_rr = []
                        seen_rr = set()
                        for rr in rr_candidates:
                            marker = id(rr)
                            if marker not in seen_rr:
                                seen_rr.add(marker)
                                unique_rr.append(rr)

                        sorted_rr = sorted(
                            unique_rr,
                            key=lambda r: (
                                r.get("modifiedon")
                                or r.get("createdon")
                                or ""
                            ),
                            reverse=True,
                        )

                        for rr in sorted_rr:
                            # Support direct fields, aliases and formatted
                            # aliases returned by the ECRM OData API.
                            req = (
                                rr.get("te_requestnumber")
                                or rr.get(
                                    "te_requestnumber@OData.Community.Display.V1.FormattedValue"
                                )
                                or rr.get("rr.te_requestnumber")
                                or rr.get(
                                    "rr.te_requestnumber@OData.Community.Display.V1.FormattedValue"
                                )
                                or rr.get("rr_x002e_te_requestnumber")
                                or rr.get(
                                    "rr_x002e_te_requestnumber@OData.Community.Display.V1.FormattedValue"
                                )
                                or rr.get("request_number")
                                or rr.get("requestnumber")
                                or ""
                            )

                            req = str(req or "").strip()

                            if req and req.lower() not in {
                                "none",
                                "null",
                                "nan",
                                "n/a",
                                "na",
                                "-",
                            }:
                                cid_ord = req
                                break

                    cid_status_rows.append({
                        "Order": str(order or "").strip(),
                        "CID": cid,
                        "Order Status": order_status,
                        "Circuit Status": circuit_status,
                        "Request Number": cid_ord,
                        "Infra status": infra_status,
                        "ESPT status": espt_status,
                        "TE note": te_note,
                    })

                # Order exists but has no Old/New CID.
                if not cid_status_rows:
                    cid_status_rows.append({
                        "Order": str(order or "").strip(),
                        "CID": "",
                        "Order Status": order_status,
                        "Circuit Status": "not found",
                        "Request Number": "",
                        "Infra status": "",
                        "ESPT status": "",
                        "TE note": "",
                    })

            cid_status_by_index[index] = cid_status_rows
            
# داخل worker لا نكتب Excel لتجنب اختلاف ترتيب النتائج
            row_values = [row.get(h, "") for h in headers]

            trans_type = str(row.get("Transmission Type") or "General").strip()
            sheet_name = "".join(c for c in trans_type if c.isalnum() or c in " _-")[:31]
            if not sheet_name:
                sheet_name = "General"

            with lock:
                completed_count += 1
                progress = (completed_count / total) * 100
                set_progress(gui, progress, f"Processed {completed_count}/{total} - {order}")
            stage(f"✓ Item {index + 1}/{total} completed", "Completed")

            if debug_mode:
                _run_smart_debug(order, context, debug_mode)

            return (index, dict(row), row_values, sheet_name)
            
            if debug_mode:
                _run_smart_debug(order, context, debug_mode)
                
        except Exception as e:
            if not isinstance(e, ExtractionCancelled):
                with lock:
                    completed_count += 1
                    set_progress(gui, (completed_count / total) * 100, f"Error on {item}")
                    error_row = {"Order": f"Error: {e}"}
                    if "Order" not in headers and headers:
                        error_row = {headers[0]: f"Error: {e}"}
                    error_values = [error_row.get(h, "") for h in headers]

                    return (index, error_row, error_values, "General")
            raise

    cancelled = False
    futures = []

    try:
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = [executor.submit(worker, i, item) for i, item in enumerate(items)]
            try:
                for future in as_completed(futures):
                    idx, row_dict, row_values, sheet_name = future.result()
                    results_by_index[idx] = (row_dict, row_values, sheet_name)
            except ExtractionCancelled:
                # Stop queued work. Running workers are allowed to finish their
                # current item so every item that actually completed is saved.
                cancelled = True
                executor.shutdown(wait=True, cancel_futures=True)

                # Collect only futures that completed successfully.
                for future in futures:
                    if not future.done() or future.cancelled():
                        continue
                    try:
                        idx, row_dict, row_values, sheet_name = future.result()
                        results_by_index[idx] = (row_dict, row_values, sheet_name)
                    except ExtractionCancelled:
                        pass
                raise

    finally:
        # Build the workbook from every successfully completed item, including
        # partial results when the user presses Cancel.
        saved_count = 0
        for i in range(total):
            if results_by_index[i]:
                _, row_values, sheet_name = results_by_index[i]
                ws.append(row_values)

                if sheet_name not in wb.sheetnames:
                    target_ws = wb.create_sheet(title=sheet_name)
                    target_ws.append(headers)
                else:
                    target_ws = wb[sheet_name]
                target_ws.append(row_values)
                saved_count += 1


        # ============================================================
        # Dedicated CID_Status sheet
        # ============================================================
        if (
            "Circuit Status" in selected
            or "Request Number" in selected
        ):
            cid_ws = wb.create_sheet(title="CID_Status")
            cid_ws.append([
                "Order", "CID", "Order Status", "Circuit Status",
                "Request Number", "Infra status", "ESPT status", "TE note", "TE note"
            ])

            # Preserve input Order sequence.
            for i in range(total):
                for cid_row in (cid_status_by_index[i] or []):
                    cid_ws.append([
                        cid_row.get("Order", ""),
                        cid_row.get("CID", ""),
                        cid_row.get("Order Status", ""),
                        cid_row.get("Circuit Status", ""),
                        cid_row.get("Request Number", ""),
                        cid_row.get("Infra status", ""),
                        cid_row.get("ESPT status", ""),
                        cid_row.get("TE note", ""),
                    ])

            cid_ws.freeze_panes = "A2"
            cid_ws.auto_filter.ref = cid_ws.dimensions
            cid_ws.column_dimensions["A"].width = 14
            cid_ws.column_dimensions["B"].width = 20
            cid_ws.column_dimensions["C"].width = 18
            cid_ws.column_dimensions["D"].width = 18
            cid_ws.column_dimensions["E"].width = 22
            cid_ws.column_dimensions["F"].width = 24
            cid_ws.column_dimensions["G"].width = 22
            cid_ws.column_dimensions["H"].width = 45

        with lock:
            wb.save(out_excel)

        if cancelled:
            gui._partial_results_count = saved_count
            print(f"Cancelled: saved {saved_count}/{total} completed items to {out_excel}")
        else:
            gui._partial_results_count = saved_count
            print(f"Final save completed: {saved_count}/{total} items -> {out_excel}")

    # إرجاع rows للـ smart_result بنفس ترتيب input
    return [results_by_index[i][0] for i in range(total) if results_by_index[i]]

def _run_smart_debug(order, ctx, debug_mode):
    """تنفيذ عمليات الحفظ في وضع الـ Debug بناءً على السياق المستلم"""
    if not debug_mode:
        return

    deps = ctx["deps"]
    session = ctx["session"]
    selected = ctx["selected"]
    order = ctx["order"]
    input_cid = ctx.get("input_cid", "")
    so = ctx.get("so")
    migration_so = ctx.get("migration_so")
    resources = ctx.get("resources")

    # تعريف البادئة المستخدمة في تسمية الملفات لضمان عملها في وضع الـ CID بدون Order
    file_prefix = order if order else input_cid

    if debug_mode: # في وضع التصحيح نحفظ البيانات الأساسية دائماً
        if ctx.get("service_orders"): save_json(file_prefix, "SERVICE_ORDERS", ctx["service_orders"])
        if so: save_json(file_prefix, "LATEST_SO", so)
        if migration_so: save_json(file_prefix, "MIGRATION_SO", migration_so)

    if resources:
        save_json(file_prefix, "RESOURCES", resources)
        if ctx.get("l3") and "Network Data" in selected: save_json(file_prefix, "L3_PARSED", ctx["l3"])

    if has_any_selected(selected, RESULT_REQUIRED_FIELDS) and ctx.get("result"):
        save_json(file_prefix, "RESULT_PARSED", ctx["result"])

    if "Current Task" in selected or "Latest Migration by E-Support SO" in selected:
        if ctx.get("tasks") or so:
            open_tasks = [t for t in (ctx.get("tasks") or []) if t.get("statecode") == 0]
            save_json(file_prefix, "TASKS_DEBUG", {"total": len(ctx.get("tasks", [])), "open": open_tasks})

        if ctx.get("migration_tasks"):
            m_tasks = ctx["migration_tasks"]
            open_m = [t for t in m_tasks if t.get("statecode") == 0 or t.get("statecode@OData.Community.Display.V1.FormattedValue") == "Open"]
            save_json(file_prefix, "MIGRATION_TASKS_DEBUG", {"total": len(m_tasks), "open": open_m})

    if ctx.get("installed"):
        save_json(file_prefix, "INSTALLED_BASE", ctx["installed"])

    if ctx.get("customer_json"):
        save_json(file_prefix, "CUSTOMER_INFO", ctx["customer_json"])
